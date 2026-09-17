#!/usr/bin/env python3
"""Cluster ADHD keyword research into intent decisions for Structuro.

Does not create one page per keyword. Every row gets a cluster, risk label
and action. Raw volume sums of variants are kept separate from representative
volume (max in cluster).

Requires: Python 3.9+ and openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from statistics import mean, median
from typing import Iterable

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CANDIDATES = [
    ROOT
    / "data"
    / "seo"
    / "export_research_nl_suggestion_keywords_history_similar_eur_2026-09_adhd.xlsx",
    Path.home()
    / "Downloads"
    / "export_research_nl_suggestion_keywords_history_similar_eur_2026-09_adhd.xlsx",
]

STOPWORDS = {
    "de",
    "het",
    "een",
    "van",
    "en",
    "voor",
    "bij",
    "met",
    "op",
    "is",
    "te",
    "in",
    "om",
    "tot",
    "uit",
    "over",
    "naar",
    "of",
    "als",
    "dan",
    "ook",
    "je",
    "ik",
    "mijn",
    "jouw",
    "zijn",
    "haar",
    "hun",
    "vs",
    "versus",
    "the",
    "a",
    "to",
}

# Tokens that flip a query into YMYL even inside an otherwise useful cluster.
YMYL_HIGH_RE = re.compile(
    r"""
    (
        diagnose|diagnostic|gediagnosticeerd|dsm|kenmerken|symptoom|symptomen|
        vragenlijst|zelftest|\btesten\b|\btest\b|\btestje\b|screening|\bchecklist\b|
        \bheb\s+ik\s+adhd\b|hoe\s+weet\s+je|of\s+je\s+adhd\s+hebt|
        medicatie|medication|medicijn|medicine|pillen|\bpil\b|vitamine|
        ritalin|concerta|methylfenidaat|dexamfetamine|
        elvanse|atomoxetine|strattera|dosering|dosis|stimulant|
        supplement|theanine|lto3|l\s*to\s*3|magnesium|ashwagandha|cbd|
        visolie|omega|lions?\s*mane|paddestoel|nootropic|
        erfelijk|erfelijke|genetica|genetisch|
        rijbewijs|cbr\b|keuring|
        behandeling|behandelen|therapie|psychiater|psycholoog|ggz|psyq|
        alcohol|drugs|wiet|cannabis|cocaine|verslaving|
        zwanger|zwangerschap|borstvoeding|
        \bkind(?:eren|je)?\b|baby|peuter|kleuter|tiener|puber|schoolkind|
        \bmeisjes?\b|\bjongens?\b|kindje|
        depressie|depressief|angststoornis|bipolar|schizofrenie
    )
    """,
    re.VERBOSE,
)
YMYL_MED_RE = re.compile(
    r"\b(autisme|ass\b|comorbid|comorbiditeit)\b"
)

ALLOW_EXISTING_DESPITE_YMYL = {
    "adhd bij vrouwen",
    "adhd vrouwen",
    "adhd vrouw",
    "vrouwen adhd",
    "vrouwen met adhd",
    "adhd bij de vrouw",
    "adhd en burn out",
    "adhd en burnout",
    "adhd burn out",
    "adhd burnout",
    "adhd en burn-out",
}

EU = "https://www.structuro.eu"


@dataclass(frozen=True)
class Seed:
    cluster_id: str
    cluster_name: str
    url: str
    page_type: str
    tier: str
    default_action: str
    default_ymyl: str
    product_fit: str
    audience_fit: str
    patterns: tuple[str, ...]
    primary_query: str


SEEDS: tuple[Seed, ...] = (
    Seed(
        "adhd_app_product",
        "ADHD-app (product/categorie)",
        f"{EU}/adhd-app/",
        "product_landing",
        "P0",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (
            r"^(adhd\s+)?app(licatie)?$",
            r"^app(licatie)?\s+adhd$",
            r"adhd\s+app(?!.*(store|storekit|ios only))",
            r"app\s+voor\s+adhd",
            r"adhd\s+apps?(?!.*beste)",
            r"executie.?app",
            r"adhd\s+hulpmiddel\s+app",
        ),
        "adhd app",
    ),
    Seed(
        "adhd_app_comparison",
        "Beste ADHD-app / vergelijking",
        f"{EU}/beste-adhd-app-nederland/",
        "comparison_guide",
        "P0",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (
            r"beste\s+adhd\s+app",
            r"adhd\s+app\s+(vergelijk|vergelijken|nederland|nl)\b",
            r"top\s+\d+\s+adhd\s+app",
            r"welke\s+adhd\s+app",
        ),
        "beste adhd app nederland",
    ),
    Seed(
        "planner_planning_agenda",
        "ADHD planner / planning / agenda",
        f"{EU}/adhd-planner-die-niet-overvraagt/",
        "guide",
        "P0",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (
            r"\bplanner\b",
            r"\bplanning\b",
            r"\bagenda\b",
            r"\bkalender\b",
            r"bullet\s+journal",
            r"\bbujo\b",
        ),
        "planning adhd",
    ),
    Seed(
        "planners_fail",
        "Waarom planners falen",
        f"{EU}/waarom-planners-falen/",
        "guide",
        "P0",
        "KEEP",
        "NONE",
        "high",
        "adults_start_friction",
        (r"planner(s)?\s+(falen|werkt\s+niet|niet\s+werken)", r"waarom\s+planner"),
        "waarom planners falen adhd",
    ),
    Seed(
        "todo_alt",
        "Alternatief voor to-do-lijst",
        f"{EU}/alternatief-voor-todo-lijst-adhd/",
        "guide",
        "P0",
        "EXPAND",
        "NONE",
        "high",
        "adults_start_friction",
        (r"to-?do", r"todo", r"takenlijst", r"taken\s+lijst"),
        "alternatief todo lijst adhd",
    ),
    Seed(
        "executive_functions",
        "Executieve functies en taakinitiatie",
        f"{EU}/executieve-functies-adhd/",
        "guide",
        "P0",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (
            r"executieve\s+funct",
            r"executive\s+func",
            r"taakinitiatie",
            r"executie(ve)?\s+dysfunct",
            r"uitvoerende\s+funct",
        ),
        "executieve functies adhd",
    ),
    Seed(
        "work",
        "ADHD op het werk",
        f"{EU}/adhd-op-het-werk/",
        "guide",
        "P0",
        "EXPAND",
        "LOW",
        "high",
        "working_adults",
        (
            r"\bwerk\b",
            r"\bkantoor\b",
            r"\bjob\b",
            r"\bbaan\b",
            r"\bcarri[eè]re\b",
            r"\bmeeting",
            r"\bcollega",
        ),
        "adhd op het werk",
    ),
    Seed(
        "start_friction",
        "Niet kunnen beginnen / taakinitiatie",
        f"{EU}/niet-kunnen-beginnen-adhd/",
        "guide",
        "P0",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (
            r"niet\s+kunnen\s+beginnen",
            r"niet\s+starten",
            r"kan\s+niet\s+beginnen",
            r"startfrictie",
            r"beginnen\s+lukt\s+niet",
            r"niet\s+aan\s+de\s+slag",
        ),
        "niet kunnen beginnen adhd",
    ),
    Seed(
        "task_paralysis",
        "Taakverlamming",
        f"{EU}/taakverlamming-adhd/",
        "guide",
        "P0",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (r"taakverlamming", r"task\s+paralysis", r"verlamd\s+door\s+taken"),
        "taakverlamming adhd",
    ),
    Seed(
        "just_start",
        "Waarom gewoon beginnen niet werkt",
        f"{EU}/waarom-gewoon-beginnen-niet-werkt/",
        "guide",
        "P0",
        "KEEP",
        "NONE",
        "high",
        "adults_start_friction",
        (r"gewoon\s+beginnen", r"gewoon\s+doen"),
        "gewoon beginnen adhd",
    ),
    Seed(
        "one_step",
        "Eén stap per dag",
        f"{EU}/een-stap-per-dag/",
        "guide",
        "P0",
        "KEEP",
        "NONE",
        "high",
        "adults_start_friction",
        (r"een\s+stap\s+per\s+dag", r"micro.?stap", r"eerste\s+stap"),
        "een stap per dag adhd",
    ),
    Seed(
        "energy",
        "Energie-first",
        f"{EU}/energie-first/",
        "guide",
        "P1",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (r"\benergie\b", r"energy\s+first"),
        "energie first adhd",
    ),
    Seed(
        "mental_load",
        "Mentale belasting",
        f"{EU}/mentale-belasting-dagstart/",
        "guide",
        "P1",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (r"mentale\s+belasting", r"mental\s+load", r"hoofd\s+vol"),
        "mentale belasting adhd",
    ),
    Seed(
        "overstimulation",
        "Overprikkeling",
        f"{EU}/overprikkeling-adhd/",
        "guide",
        "P1",
        "EXPAND",
        "LOW",
        "medium",
        "adults_start_friction",
        (r"overprikkel", r"overstimul", r"prikkelarm", r"shutdown"),
        "overprikkeling adhd",
    ),
    Seed(
        "procrastination",
        "Uitstelgedrag",
        f"{EU}/adhd-uitstelgedrag/",
        "guide",
        "P1",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (r"uitstel", r"procrastin"),
        "adhd uitstelgedrag",
    ),
    Seed(
        "women_practical",
        "ADHD bij vrouwen (executie, geen diagnose)",
        f"{EU}/adhd-bij-vrouwen/",
        "guide",
        "P1",
        "EXPAND",
        "MEDIUM",
        "medium",
        "adult_women",
        (r"vrouw(en)?", r"vrouwelijk"),
        "adhd bij vrouwen",
    ),
    Seed(
        "focus_no_streaks",
        "Focus zonder streaks",
        f"{EU}/adhd-focus-zonder-streaks/",
        "guide",
        "P1",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (r"\bfocus\b", r"\bconcentratie\b", r"\baandacht\b", r"streak", r"pomodoro"),
        "adhd focus",
    ),
    Seed(
        "time_blindness",
        "Tijdblindheid",
        f"{EU}/tijdblindheid-adhd/",
        "guide",
        "P1",
        "EXPAND",
        "LOW",
        "medium",
        "adults_start_friction",
        (r"tijdblind", r"time\s+blind", r"tijdsinschatting", r"tijd\s+zien"),
        "tijdblindheid adhd",
    ),
    Seed(
        "decision",
        "Keuzestress",
        f"{EU}/adhd-keuzestress/",
        "guide",
        "P1",
        "EXPAND",
        "LOW",
        "high",
        "adults_start_friction",
        (r"keuzestress", r"keuze.?overload", r"besluiteloos", r"niet\s+kunnen\s+kiezen"),
        "adhd keuzestress",
    ),
    Seed(
        "long_list",
        "Takenlijst te lang",
        f"{EU}/takenlijst-te-lang-adhd/",
        "guide",
        "P1",
        "EXPAND",
        "NONE",
        "high",
        "adults_start_friction",
        (r"te\s+lang(e)?\s+lijst", r"overweldig(d|ing).*(lijst|taken)", r"inbox\s+nul"),
        "takenlijst te lang adhd",
    ),
    Seed(
        "morning",
        "Ochtendroutine",
        f"{EU}/adhd-ochtendroutine/",
        "guide",
        "P1",
        "EXPAND",
        "LOW",
        "medium",
        "adults_start_friction",
        (r"ochtend", r"morning\s+routine"),
        "adhd ochtendroutine",
    ),
    Seed(
        "burnout",
        "ADHD en burn-out (energie, geen behandeling)",
        f"{EU}/adhd-en-burn-out/",
        "guide",
        "P1",
        "EXPAND",
        "MEDIUM",
        "medium",
        "adults_start_friction",
        (r"burn.?out", r"opgebrand", r"leegte"),
        "adhd en burn-out",
    ),
    Seed(
        "cycle",
        "Cyclus (optionele productcontext)",
        f"{EU}/cyclus/",
        "product_context",
        "P1",
        "KEEP",
        "MEDIUM",
        "medium",
        "adult_women",
        (r"\bcyclus\b", r"hormoon", r"menstru"),
        "adhd cyclus",
    ),
    Seed(
        "tiimo",
        "Structuro of Tiimo",
        f"{EU}/structuro-of-tiimo/",
        "comparison",
        "P0",
        "KEEP",
        "NONE",
        "high",
        "switchers",
        (r"\btiimo\b",),
        "structuro of tiimo",
    ),
    Seed(
        "structured",
        "Structuro of Structured",
        f"{EU}/structuro-of-structured/",
        "comparison",
        "P0",
        "KEEP",
        "NONE",
        "high",
        "switchers",
        (r"\bstructured\b",),
        "structuro of structured",
    ),
    Seed(
        "todoist",
        "Structuro of Todoist",
        f"{EU}/structuro-of-todoist/",
        "comparison",
        "P0",
        "KEEP",
        "NONE",
        "high",
        "switchers",
        (r"\btodoist\b",),
        "structuro of todoist",
    ),
    Seed(
        "research",
        "Eigen onderzoek taakinitiatie",
        f"{EU}/onderzoek/",
        "research",
        "P0",
        "KEEP",
        "LOW",
        "high",
        "press_and_seo",
        (r"structuro.*onderzoek", r"gebruikersonderzoek", r"taakinitiatie.*onderzoek"),
        "structuro onderzoek",
    ),
    Seed(
        "brand",
        "Merk / homepage",
        f"{EU}/",
        "homepage",
        "P0",
        "KEEP",
        "NONE",
        "high",
        "brand",
        (r"^structuro$", r"structuro\s+app", r"structuro\s+adhd"),
        "structuro",
    ),
)

COMPILED_SEEDS = tuple(
    (seed, tuple(re.compile(p) for p in seed.patterns)) for seed in SEEDS
)

IRRELEVANT_RE = re.compile(
    r"""
    \b(
        celebrity|bekend|famous|youtube\s+ster|tiktok\s+ster|
        meme|grappig|joke|t-shirt|hoodie|tattoo|
        harry\s+potter|marvel|spotify\s+playlist\s+only
    )\b
    """,
    re.VERBOSE,
)


def find_workbook(explicit: str | None) -> Path:
    if explicit:
        path = Path(explicit).expanduser().resolve()
        if not path.is_file():
            raise SystemExit(f"Workbook not found: {path}")
        return path
    fuller: list[Path] = []
    for folder in (ROOT / "data" / "seo", Path.home() / "Downloads", ROOT):
        if not folder.is_dir():
            continue
        fuller.extend(folder.glob("*suggestion_keywords*.xlsx"))
        fuller.extend(folder.glob("*export_research*adhd*.xlsx"))
    unique = []
    seen = set()
    for path in fuller + DEFAULT_CANDIDATES:
        key = str(path)
        if key in seen or not path.is_file():
            continue
        seen.add(key)
        unique.append(path)
    if not unique:
        raise SystemExit(
            "No keyword workbook found. Pass --input path/to.xlsx"
        )
    # Prefer the largest file in case a fuller export appears later.
    unique.sort(key=lambda p: p.stat().st_size, reverse=True)
    return unique[0]


def load_rows(path: Path) -> tuple[list[str], list[dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("openpyxl is required: python3 -m pip install openpyxl") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        wb.close()
        raise SystemExit("Workbook has no header row")
    header = [str(c).strip() if c is not None else "" for c in header_row]
    wanted = {
        "Zoekwoord": "original_keyword",
        "Moeilijkheid": "difficulty",
        "Zoek Vol.": "volume",
        "Zoekintentie": "intent",
        "SERP Kenmerken": "serp_features",
        "CPC": "cpc",
        "Concurrentie": "competition",
    }
    index = {}
    for key, dest in wanted.items():
        if key not in header:
            wb.close()
            raise SystemExit(f"Missing column {key!r}. Found: {header}")
        index[dest] = header.index(key)

    out = []
    for i, raw in enumerate(rows_iter, start=2):
        if raw is None:
            continue
        values = list(raw)
        # Pad short rows instead of dropping them.
        if len(values) < len(header):
            values.extend([None] * (len(header) - len(values)))
        kw = values[index["original_keyword"]]
        if kw is None or str(kw).strip() == "":
            continue
        out.append(
            {
                "excel_row": i,
                "original_keyword": str(kw).strip(),
                "difficulty": _num(values[index["difficulty"]]),
                "volume": _num(values[index["volume"]]),
                "intent": _str(values[index["intent"]]),
                "serp_features": _str(values[index["serp_features"]]),
                "cpc": _num(values[index["cpc"]]),
                "competition": _num(values[index["competition"]]),
            }
        )
    wb.close()
    return header, out


def _num(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def _str(value) -> str:
    return "" if value is None else str(value).strip()


def normalize_keyword(keyword: str) -> str:
    text = unicodedata.normalize("NFKC", keyword).lower().strip()
    text = text.replace("’", "'").replace("`", "'")
    text = text.replace("&", " en ")
    text = re.sub(r"[\u2010-\u2015\-]+", " ", text)
    text = re.sub(r"[^\w\s+']", " ", text, flags=re.UNICODE)
    text = re.sub(r"\s+", " ", text).strip()
    text = text.replace("l theanine", "l-theanine")
    text = text.replace("burn out", "burn-out")
    return text


def token_signature(normalized: str) -> str:
    tokens = [t for t in normalized.replace("-", " ").split() if t and t not in STOPWORDS]
    if not tokens:
        tokens = normalized.split()
    return " ".join(sorted(tokens))


def ymyl_risk(normalized: str) -> str:
    if normalized in ALLOW_EXISTING_DESPITE_YMYL:
        return "MEDIUM"
    if YMYL_HIGH_RE.search(normalized):
        return "HIGH"
    if YMYL_MED_RE.search(normalized):
        return "MEDIUM"
    return "NONE"


def match_seed(normalized: str) -> Seed | None:
    # Comparison and brand before generic app.
    priority_ids = (
        "adhd_app_comparison",
        "tiimo",
        "structured",
        "todoist",
        "brand",
        "planners_fail",
    )
    for seed, crex in COMPILED_SEEDS:
        if seed.cluster_id not in priority_ids:
            continue
        if any(rx.search(normalized) for rx in crex):
            return seed
    for seed, crex in COMPILED_SEEDS:
        if seed.cluster_id in priority_ids:
            continue
        if any(rx.search(normalized) for rx in crex):
            return seed
    return None


def is_app_query(normalized: str) -> bool:
    return bool(re.search(r"\bapp(s|licatie)?\b", normalized))


def classify_row(row: dict) -> dict:
    original = row["original_keyword"]
    normalized = normalize_keyword(original)
    signature = token_signature(normalized)
    risk = ymyl_risk(normalized)
    seed = match_seed(normalized)

    # App + planner should sit with the product app cluster when "app" is present.
    if seed and seed.cluster_id == "planner_planning_agenda" and is_app_query(normalized):
        seed = next(s for s in SEEDS if s.cluster_id == "adhd_app_product")

    if risk == "HIGH" and normalized not in ALLOW_EXISTING_DESPITE_YMYL:
        womenish = bool(re.search(r"vrouw", normalized))
        action = "HOLD_YMYL"
        cluster_id = "hold_ymyl"
        cluster_name = "YMYL / niet targeten zonder medische redactie"
        url = ""
        page_type = "none"
        tier = "P3"
        product_fit = "blocked"
        audience_fit = "ymyl"
        primary = False
        reason = (
            "Hoog YMYL-risico (diagnose, medicatie, supplement, kind, erfelijkheid "
            "of juridisch). Lage KD of hoog volume weegt hier niet tegenop."
        )
        if womenish and re.search(r"kenmerk|symptoom|diagnose|test|medicatie", normalized):
            reason = (
                "Volume rond vrouwen is relevant, maar deze query is diagnostisch. "
                "Bestaande pagina blijft praktisch; deze term niet als funnel gebruiken."
            )
        return _pack(
            row,
            normalized,
            signature,
            cluster_id,
            cluster_name,
            url,
            action,
            "supporting",
            audience_fit,
            product_fit,
            risk,
            page_type,
            tier,
            reason,
        )

    if IRRELEVANT_RE.search(normalized):
        return _pack(
            row,
            normalized,
            signature,
            "irrelevant",
            "Niet relevant / ruis",
            "",
            "IGNORE_IRRELEVANT",
            "supporting",
            "none",
            "none",
            risk if risk != "NONE" else "NONE",
            "none",
            "P3",
            "Geen product- of topicfit voor Structuro.",
        )

    if seed:
        action = seed.default_action
        role = "primary" if normalized == normalize_keyword(seed.primary_query) else "supporting"
        if action == "NEW_PAGE" and role == "supporting":
            action = "SUPPORTING_KEYWORD"
        elif action in {"EXPAND", "KEEP", "UPDATE"} and role == "supporting":
            action = "SUPPORTING_KEYWORD"
        # adhd agenda: shopping (fysieke agenda) en executie overlappen in NL-SERP.
        # Geen aparte URL. De planner-gids behandelt het verschil expliciet.
        reason = _seed_reason(seed, action, risk)
        return _pack(
            row,
            normalized,
            signature,
            seed.cluster_id,
            seed.cluster_name,
            seed.url,
            action,
            role,
            seed.audience_fit,
            seed.product_fit,
            seed.default_ymyl if risk == "NONE" else risk,
            seed.page_type,
            seed.tier,
            reason,
        )

    # Generic head term.
    if normalized == "adhd":
        return _pack(
            row,
            normalized,
            signature,
            "generic_adhd",
            "Generieke head term ADHD",
            f"{EU}/",
            "IGNORE_IRRELEVANT",
            "supporting",
            "everyone",
            "weak",
            "MEDIUM",
            "none",
            "P3",
            "Te breed en YMYL-aangrenzend. Geen eigen URL; homepage is merk, geen encyclopedie.",
        )

    # Leftover: signature cluster, no new URL by default.
    return _pack(
        row,
        normalized,
        signature,
        f"residual:{signature}",
        f"Restcluster: {signature}",
        "",
        "IGNORE_IRRELEVANT",
        "supporting",
        "unknown",
        "low",
        risk,
        "none",
        "P2",
        "Geen eigen pagina. Alleen publiceren bij een écht onderscheidende informatiebehoefte.",
    )


def _seed_reason(seed: Seed, action: str, risk: str) -> str:
    if action == "NEW_PAGE":
        return (
            f"Gerechtvaardigde nieuwe canonical voor {seed.cluster_name}. "
            "Varianten hangen als supporting keywords aan dezelfde URL."
        )
    if action == "NEEDS_SERP_REVIEW":
        return (
            "Mogelijke SERP-split (fysieke agenda vs software). Default: bestaande "
            "planner-gids, geen extra URL tot de SERP materieel anders is."
        )
    if action == "SUPPORTING_KEYWORD":
        return f"Synoniem of woordvolgorde-variant van {seed.cluster_name}; geen extra URL."
    if seed.default_action in {"EXPAND", "UPDATE"}:
        return f"Bestaande canonical versterken: {seed.url}"
    return f"Bestaande canonical behouden: {seed.url}"


def _pack(
    row: dict,
    normalized: str,
    signature: str,
    cluster_id: str,
    cluster_name: str,
    url: str,
    action: str,
    role: str,
    audience: str,
    product: str,
    risk: str,
    page_type: str,
    tier: str,
    reason: str,
) -> dict:
    volume = row["volume"]
    difficulty = row["difficulty"]
    fit = {"high": 1.0, "medium": 0.6, "low": 0.25, "weak": 0.1, "none": 0.05, "blocked": 0.02}.get(
        product, 0.2
    )
    tier_w = {"P0": 1.0, "P1": 0.55, "P2": 0.2, "P3": 0.08}[tier]
    ymyl_w = {"NONE": 1.0, "LOW": 0.9, "MEDIUM": 0.5, "HIGH": 0.05}.get(risk, 0.7)
    intent = row["intent"].upper()
    intent_w = 1.15 if "C" in intent else 1.0 if intent.startswith("I") else 0.85
    score = round((volume * fit * tier_w * ymyl_w * intent_w) / (1.0 + difficulty), 3)
    business = "high" if tier == "P0" and risk in {"NONE", "LOW"} else (
        "medium" if tier in {"P0", "P1"} else "low"
    )
    return {
        **row,
        "normalized_keyword": normalized,
        "token_signature": signature,
        "cluster_id": cluster_id,
        "cluster_name": cluster_name,
        "primary_or_supporting": role,
        "audience_fit": audience,
        "product_fit": product,
        "business_value": business,
        "ymyl_risk": risk,
        "existing_url": "" if action == "NEW_PAGE" else url,
        "recommended_url": url,
        "recommended_action": action,
        "page_type": page_type,
        "priority_tier": tier,
        "reason": reason,
        "opportunity_score_heuristic": score,
        "ai_overview_mentioned": "ai-overzicht" in row["serp_features"].lower(),
        "paa_mentioned": "mensen vragen ook" in row["serp_features"].lower(),
    }


def assign_cluster_volume_fields(rows: list[dict]) -> None:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[row["cluster_id"]].append(row)
    for cluster_id, items in grouped.items():
        vols = [r["volume"] for r in items]
        raw_sum = sum(vols)
        representative = max(vols) if vols else 0
        # Pick one primary per cluster when missing.
        primaries = [r for r in items if r["primary_or_supporting"] == "primary"]
        if not primaries:
            best = max(items, key=lambda r: (r["volume"], -r["difficulty"]))
            best["primary_or_supporting"] = "primary"
            if best["recommended_action"] == "SUPPORTING_KEYWORD":
                seed_action = next(
                    (s.default_action for s in SEEDS if s.cluster_id == cluster_id),
                    best["recommended_action"],
                )
                if seed_action in {"EXPAND", "UPDATE", "KEEP", "NEW_PAGE"}:
                    best["recommended_action"] = seed_action
        for row in items:
            row["cluster_variant_count"] = len(items)
            row["cluster_volume_representative"] = representative
            row["cluster_volume_sum_raw_not_deduped"] = raw_sum


def opportunity_rank(rows: list[dict]) -> list[dict]:
    scored = []
    seen = set()
    for row in sorted(rows, key=lambda r: r["opportunity_score_heuristic"], reverse=True):
        cid = row["cluster_id"]
        if cid in seen:
            continue
        if row["recommended_action"] in {"HOLD_YMYL", "IGNORE_IRRELEVANT"}:
            continue
        seen.add(cid)
        scored.append(row)
    return scored


FIELDNAMES = [
    "excel_row",
    "original_keyword",
    "normalized_keyword",
    "difficulty",
    "volume",
    "intent",
    "serp_features",
    "cpc",
    "competition",
    "cluster_id",
    "cluster_name",
    "primary_or_supporting",
    "audience_fit",
    "product_fit",
    "business_value",
    "ymyl_risk",
    "existing_url",
    "recommended_url",
    "recommended_action",
    "page_type",
    "priority_tier",
    "reason",
    "cluster_variant_count",
    "cluster_volume_representative",
    "cluster_volume_sum_raw_not_deduped",
    "opportunity_score_heuristic",
    "ai_overview_mentioned",
    "paa_mentioned",
]


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDNAMES, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def write_page_map(path: Path, rows: list[dict]) -> None:
    clusters: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        clusters[row["cluster_id"]].append(row)
    fieldnames = [
        "cluster_id",
        "cluster_name",
        "priority_tier",
        "ymyl_risk_max",
        "recommended_url",
        "primary_keyword",
        "keyword_count",
        "volume_representative",
        "volume_sum_raw_not_deduped",
        "recommended_action_primary",
        "page_type",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for cid, items in sorted(
            clusters.items(),
            key=lambda kv: -max(r["volume"] for r in kv[1]),
        ):
            primary = max(items, key=lambda r: (r["primary_or_supporting"] == "primary", r["volume"]))
            risks = {r["ymyl_risk"] for r in items}
            order = ["HIGH", "MEDIUM", "LOW", "NONE"]
            risk_max = next(r for r in order if r in risks)
            writer.writerow(
                {
                    "cluster_id": cid,
                    "cluster_name": primary["cluster_name"],
                    "priority_tier": primary["priority_tier"],
                    "ymyl_risk_max": risk_max,
                    "recommended_url": primary["recommended_url"],
                    "primary_keyword": next(
                        (r["original_keyword"] for r in items if r["primary_or_supporting"] == "primary"),
                        primary["original_keyword"],
                    ),
                    "keyword_count": len(items),
                    "volume_representative": max(r["volume"] for r in items),
                    "volume_sum_raw_not_deduped": sum(r["volume"] for r in items),
                    "recommended_action_primary": primary["recommended_action"]
                    if primary["primary_or_supporting"] == "primary"
                    else items[0]["recommended_action"],
                    "page_type": primary["page_type"],
                }
            )


def summarize(path: Path, rows: list[dict]) -> dict:
    kd30 = [r for r in rows if r["difficulty"] < 30]
    actions = defaultdict(int)
    for row in rows:
        actions[row["recommended_action"]] += 1
    mapped_existing = sum(1 for r in rows if r["recommended_url"] and r["recommended_action"] not in {"NEW_PAGE", "HOLD_YMYL", "IGNORE_IRRELEVANT"})
    new_page = sum(1 for r in rows if r["recommended_action"] == "NEW_PAGE")
    supporting = sum(1 for r in rows if r["recommended_action"] == "SUPPORTING_KEYWORD")
    hold = sum(1 for r in rows if r["recommended_action"] == "HOLD_YMYL")
    irrelevant = sum(1 for r in rows if r["recommended_action"] == "IGNORE_IRRELEVANT")
    clusters = {r["cluster_id"] for r in rows}
    semantic_clusters = {c for c in clusters if not c.startswith("residual:")}
    residual = {c for c in clusters if c.startswith("residual:")}
    summary = {
        "workbook": str(path),
        "row_count": len(rows),
        "kd_lt_30": len(kd30),
        "kd_lt_30_share": round(len(kd30) / len(rows), 4) if rows else 0,
        "volume_sum_kd_lt_30_raw_not_deduped": sum(r["volume"] for r in kd30),
        "volume_median_kd_lt_30": median([r["volume"] for r in kd30]) if kd30 else 0,
        "volume_mean_kd_lt_30": round(mean([r["volume"] for r in kd30]), 1) if kd30 else 0,
        "kd_lt_30_volume_ge_50": sum(1 for r in kd30 if r["volume"] >= 50),
        "kd_lt_30_volume_ge_90": sum(1 for r in kd30 if r["volume"] >= 90),
        "kd_lt_30_volume_ge_210": sum(1 for r in kd30 if r["volume"] >= 210),
        "kd_lt_30_volume_ge_300": sum(1 for r in kd30 if r["volume"] >= 300),
        "kd_lt_30_volume_ge_480": sum(1 for r in kd30 if r["volume"] >= 480),
        "kd_lt_30_volume_ge_590": sum(1 for r in kd30 if r["volume"] >= 590),
        "intent_I": sum(1 for r in rows if r["intent"].upper().startswith("I")),
        "intent_blank": sum(1 for r in rows if not r["intent"]),
        "intent_LC": sum(1 for r in rows if "L" in r["intent"].upper() and "C" in r["intent"].upper()),
        "ai_overview_rows": sum(1 for r in rows if r["ai_overview_mentioned"]),
        "paa_rows": sum(1 for r in rows if r["paa_mentioned"]),
        "cluster_count_including_residual": len(clusters),
        "semantic_seed_or_hold_clusters": len(semantic_clusters),
        "residual_signature_clusters": len(residual),
        "action_counts": dict(actions),
        "assigned_existing_or_expand": mapped_existing,
        "assigned_new_page_rows": new_page,
        "assigned_supporting": supporting,
        "assigned_hold_ymyl": hold,
        "assigned_irrelevant": irrelevant,
        "unexplained_orphans": sum(1 for r in rows if not r["recommended_action"]),
    }
    return summary


def self_test() -> None:
    samples = {
        "planning adhd": "planner_planning_agenda",
        "adhd planning": "planner_planning_agenda",
        "planner adhd": "planner_planning_agenda",
        "adhd planner": "planner_planning_agenda",
        "adhd agenda": "planner_planning_agenda",
        "adhd app": "adhd_app_product",
        "adhd planner app": "adhd_app_product",
        "beste adhd app nederland": "adhd_app_comparison",
        "executieve functies adhd": "executive_functions",
        "adhd werk": "work",
        "adhd op het werk": "work",
        "l-theanine adhd": None,
        "adhd diagnose": None,
        "meisjes met adhd": None,
        "adhd bij vrouwen": "women_practical",
        "adhd testen": None,
        "adhd pillen": None,
    }
    failed = []
    agenda = classify_row(
        {
            "excel_row": 0,
            "original_keyword": "adhd agenda",
            "difficulty": 21,
            "volume": 210,
            "intent": "C",
            "serp_features": "",
            "cpc": 0,
            "competition": 0,
        }
    )
    if agenda["recommended_action"] != "SUPPORTING_KEYWORD":
        failed.append(
            (
                "adhd agenda",
                "expected SUPPORTING_KEYWORD",
                agenda["recommended_action"],
                agenda["cluster_id"],
            )
        )
    hyper = classify_row(
        {
            "excel_row": 0,
            "original_keyword": "hyperfocusing adhd",
            "difficulty": 5,
            "volume": 590,
            "intent": "I",
            "serp_features": "",
            "cpc": 0,
            "competition": 0,
        }
    )
    if hyper["cluster_id"] == "focus_no_streaks":
        failed.append(
            (
                "hyperfocusing adhd",
                "must not map to focus_no_streaks",
                hyper["cluster_id"],
                hyper["recommended_action"],
            )
        )
    for kw, expect in samples.items():
        row = classify_row(
            {
                "excel_row": 0,
                "original_keyword": kw,
                "difficulty": 10,
                "volume": 100,
                "intent": "I",
                "serp_features": "",
                "cpc": 0,
                "competition": 0,
            }
        )
        cid = row["cluster_id"]
        if expect is None:
            if row["recommended_action"] != "HOLD_YMYL":
                failed.append((kw, "expected HOLD_YMYL", row["recommended_action"], cid))
        elif cid != expect:
            failed.append((kw, expect, cid, row["recommended_action"]))
    assert normalize_keyword("L-Theanine  ADHD") == "l-theanine adhd"
    assert token_signature("adhd op het werk") == token_signature("adhd werk")
    if failed:
        raise SystemExit(f"self-test failed: {failed}")
    print("self-test ok")


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Cluster Structuro keyword export")
    parser.add_argument("--input", help="Path to xlsx")
    parser.add_argument(
        "--out-dir",
        default=str(ROOT / "docs" / "seo"),
        help="Directory for CSV output",
    )
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args(list(argv) if argv is not None else None)
    if args.self_test:
        self_test()
        return 0

    workbook = find_workbook(args.input)
    header, raw_rows = load_rows(workbook)
    print(f"workbook\t{workbook}")
    print(f"columns\t{header}")
    print(f"row_count\t{len(raw_rows)}")
    classified = [classify_row(r) for r in raw_rows]
    assign_cluster_volume_fields(classified)
    out_dir = Path(args.out_dir)
    write_csv(out_dir / "keyword-map.csv", classified)
    write_page_map(out_dir / "page-keyword-map.csv", classified)
    summary = summarize(workbook, classified)
    (out_dir / "keyword-summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    print("\nTop opportunity clusters (heuristic, not a ranking prediction):")
    for row in opportunity_rank(classified)[:20]:
        print(
            f"- {row['cluster_name']}: {row['original_keyword']} "
            f"vol={row['cluster_volume_representative']:.0f} kd={row['difficulty']:.0f} "
            f"{row['recommended_action']} {row['recommended_url']}"
        )
    if summary["unexplained_orphans"]:
        raise SystemExit("Some keywords have no action")
    return 0


if __name__ == "__main__":
    sys.exit(main())
